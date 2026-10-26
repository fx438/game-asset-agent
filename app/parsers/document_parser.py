"""需求文档解析：把 Word / Excel / PDF / 文本统一成"正文 + 表格"两部分。

表格单独保留，是因为规范的资源清单表可以直接结构化读取，不必交给大模型猜。
"""
import csv
import io
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_SUFFIXES = {".docx", ".xlsx", ".pdf", ".txt", ".md", ".csv"}
MAX_TEXT_CHARS = 30000


@dataclass
class ParsedDocument:
    text: str
    tables: list[list[list[str]]] = field(default_factory=list)

    def as_prompt_text(self) -> str:
        """拼成给大模型看的文本，表格转成 Markdown 风格的行。"""
        parts = [self.text.strip()]
        for i, table in enumerate(self.tables, 1):
            rows = [" | ".join(cell for cell in row) for row in table if any(row)]
            parts.append(f"[表格{i}]\n" + "\n".join(rows))
        return "\n\n".join(p for p in parts if p)[:MAX_TEXT_CHARS]


class UnsupportedDocument(ValueError):
    pass


def _cell(value) -> str:
    return "" if value is None else str(value).strip()


def _parse_docx(data: bytes) -> ParsedDocument:
    from docx import Document

    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    tables = [[[_cell(c.text) for c in row.cells] for row in t.rows] for t in doc.tables]
    return ParsedDocument(text, tables)


def _parse_xlsx(data: bytes) -> ParsedDocument:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    tables = []
    for ws in wb.worksheets:
        rows = [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if rows:
            tables.append(rows)
    wb.close()
    return ParsedDocument("", tables)


def _parse_pdf(data: bytes) -> ParsedDocument:
    import pdfplumber

    texts, tables = [], []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            texts.append(page.extract_text() or "")
            for t in page.extract_tables():
                tables.append([[_cell(c) for c in row] for row in t])
    return ParsedDocument("\n".join(texts), tables)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "gbk"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        return _parse_docx(data)
    if suffix == ".xlsx":
        return _parse_xlsx(data)
    if suffix == ".pdf":
        return _parse_pdf(data)
    if suffix == ".csv":
        return ParsedDocument("", [[_cell(c) for c in row] for row in csv.reader(io.StringIO(_decode(data)))])
    if suffix in (".txt", ".md"):
        return ParsedDocument(_decode(data))
    raise UnsupportedDocument(f"不支持的文件类型 {suffix}，支持：{', '.join(sorted(SUPPORTED_SUFFIXES))}")
