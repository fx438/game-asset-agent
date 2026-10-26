import io

import pytest

from app.agents import prompt_agent, requirement_agent, review_agent
from app.agents.styles import get_style
from app.models import ProjectStatus as PS
from app.parsers.document_parser import ParsedDocument, UnsupportedDocument, parse_document
from app.state_machine import InvalidTransition, can_transition
from tests.conftest import make_xlsx


def test_state_machine_rejects_illegal_transition():
    assert can_transition(PS.MASTER_CONFIRMING, PS.GENERATING)
    assert can_transition(PS.HUMAN_REVIEWING, PS.GENERATING)       # 人工驳回回流
    assert not can_transition(PS.PLAN_CONFIRMING, PS.GENERATING)   # 不能跳过风格母版
    assert not can_transition(PS.ARCHIVED, PS.GENERATING)


def test_compare_and_set_only_one_winner(client):
    from app.db import session_scope
    from app.models import Project
    from app.state_machine import compare_and_set, transition

    with session_scope() as s:
        p = Project(name="cas", source_filename="x.txt", source_key="k", status=PS.GENERATING)
        s.add(p)
        s.flush()
        pid = p.id
    with session_scope() as s:
        assert compare_and_set(s, pid, PS.GENERATING, PS.HUMAN_REVIEWING) is True
    with session_scope() as s:
        assert compare_and_set(s, pid, PS.GENERATING, PS.HUMAN_REVIEWING) is False
    with session_scope() as s, pytest.raises(InvalidTransition):
        transition(s, s.get(Project, pid), PS.PARSING)


def test_table_extraction_cleans_data():
    data = make_xlsx([
        ["资源清单"],
        ["英雄", "职业", "动作", "数量", "备注"],
        ["亚瑟", "战士", "攻击", "6张", "持大剑"],
        ["亚瑟", "战士", "攻击", "4", ""],          # 重复，合并取较大值
        ["妲己", "法师", "待机", "99", ""],          # 超上限，截断
        ["小乔", "法师", "受击", "若干", ""],        # 数量无法识别，跳过
    ])
    doc = parse_document("list.xlsx", data)
    result = requirement_agent.run(doc)
    assert result.method == "table"
    tasks = {(t.hero, t.action): t for t in result.tasks}
    assert tasks[("亚瑟", "攻击")].count == 6
    assert tasks[("亚瑟", "攻击")].description == "持大剑"
    assert tasks[("妲己", "待机")].count == 16
    assert ("小乔", "受击") not in tasks
    assert any("重复" in i for i in result.issues)
    assert any("超过上限" in i for i in result.issues)
    assert any("无法识别" in i for i in result.issues)


def test_free_text_extraction_in_mock_mode():
    doc = parse_document("req.txt", "需求如下：\n亚瑟 战士 攻击 4张\n妲己 法师 待机 2张\n".encode())
    result = requirement_agent.run(doc)
    assert result.method == "mock"
    assert [(t.hero, t.role, t.action, t.count) for t in result.tasks] == [("亚瑟", "战士", "攻击", 4), ("妲己", "法师", "待机", 2)]


def test_docx_table_is_parsed():
    from docx import Document

    d = Document()
    d.add_paragraph("第一期角色资源")
    table = d.add_table(rows=2, cols=4)
    for cell, text in zip(table.rows[0].cells, ["英雄", "角色", "动作", "帧数"]):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ["亚瑟", "战士", "行走", "6"]):
        cell.text = text
    buf = io.BytesIO()
    d.save(buf)
    doc = parse_document("req.docx", buf.getvalue())
    assert "第一期角色资源" in doc.text
    assert requirement_agent.run(doc).tasks[0].count == 6


def test_unsupported_file_type():
    with pytest.raises(UnsupportedDocument):
        parse_document("a.exe", b"x")


def test_prompt_keeps_style_prefix_identical():
    style = get_style("pixel")
    frames = prompt_agent.expand_frames("攻击", 3, "侧面", "")
    prompts = [prompt_agent.compose_prompt(style, "冷色调", "红发少年", "亚瑟", "战士", f, "侧面") for f in frames]
    assert len(prompts) == 3
    assert all(p.startswith(style.prefix + "，冷色调") for p in prompts)


def test_rule_check_detects_broken_and_wrong_size():
    from PIL import Image

    assert review_agent.rule_check(b"not an image") == ["图片文件损坏或格式无法识别"]
    buf = io.BytesIO()
    Image.new("RGB", (100, 100)).save(buf, format="PNG")
    assert "不符" in review_agent.rule_check(buf.getvalue())[0]


def test_empty_document_raises():
    from app.ai.errors import AIError

    with pytest.raises(AIError):
        requirement_agent.run(ParsedDocument("没有任何任务"))
